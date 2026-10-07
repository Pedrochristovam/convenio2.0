"""
Documentos da diligência da GECOV: Anexo – Solicitação de Documentos (modelo
oficial em paisagem) e minuta do Ofício de solicitação de documentos
complementares. Ambos são montados a partir de resultado["parecer"].
"""

import io
import re
from copy import deepcopy
from datetime import date
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from backend.analise_documental.minuta_ad import MODELOS_DIR
from backend.analise_documental.texto import fmt_money
from backend.analise_documental.varredura import data_extenso

LETRAS = "abcdefghijklmnopqrstuvwxyz"


def _sem_parecer() -> ValueError:
    return ValueError("Dossiê sem parecer no formato da GECOV: use 'Recalcular' para atualizar a análise.")


def _texto(p: Paragraph, texto: str, destacar: bool = False, negrito: Optional[bool] = None):
    runs = p.runs
    if runs:
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        run = runs[0]
        run.text = texto
    else:
        run = p.add_run(texto)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW if destacar else None
    if negrito is not None:
        run.bold = negrito
    return run


def _sem_numeracao(p: Paragraph):
    ppr = p._p.find(qn("w:pPr"))
    if ppr is not None:
        for num in ppr.findall(qn("w:numPr")):
            ppr.remove(num)


def _item_lista(p: Paragraph, letra: str, texto: str):
    _sem_numeracao(p)
    _texto(p, f"{letra}) {texto}", negrito=False)
    p.paragraph_format.left_indent = Cm(0.63)
    p.paragraph_format.first_line_indent = Cm(-0.63)


def _sub_linha(p: Paragraph, texto: str):
    _sem_numeracao(p)
    _texto(p, texto, negrito=False)
    p.paragraph_format.left_indent = Cm(0.63)
    p.paragraph_format.first_line_indent = Cm(0)


def _preencher_celula_lista(cell, titulo: str, itens: List[Any]):
    """Célula com título, linha em branco e itens a), b)… (itens dict têm subitens sem letra)."""
    paras = cell.paragraphs
    modelo = next((q for q in paras if q._p.find(".//" + qn("w:numPr")) is not None), paras[-1])
    modelo = deepcopy(modelo._p)
    _texto(paras[0], titulo)
    for q in paras[2:]:
        q._p.getparent().remove(q._p)
    if len(paras) < 2:
        cell.add_paragraph("")
    ultimo = cell.paragraphs[1]
    for i, item in enumerate(itens):
        texto, subs = (item["texto"], item.get("sub") or []) if isinstance(item, dict) else (item, [])
        novo = Paragraph(deepcopy(modelo), cell)
        ultimo._p.addnext(novo._p)
        _item_lista(novo, LETRAS[i % 26], texto)
        ultimo = novo
        for s in subs:
            sub = Paragraph(deepcopy(modelo), cell)
            ultimo._p.addnext(sub._p)
            _sub_linha(sub, s)
            ultimo = sub


def _celulas_unicas(row) -> List:
    out = []
    for c in row.cells:
        if not out or c._tc is not out[-1]._tc:
            out.append(c)
    return out


def gerar_solicitacao(resultado: Dict[str, Any]) -> Tuple[bytes, str]:
    parecer = resultado.get("parecer")
    if not parecer:
        raise _sem_parecer()
    itens = parecer.get("solicitacao") or []
    if not itens:
        raise ValueError("Não há documentos a solicitar: a análise não apontou pendências para diligência.")
    d = parecer["dados"]
    doc = Document(str(MODELOS_DIR / "solicitacao.docx"))
    tabela = doc.tables[0]
    rows = list(tabela.rows)

    _texto(_celulas_unicas(rows[1])[0].paragraphs[0], f"Convenente: Mun. {d.get('municipio') or '---'}")
    p = _celulas_unicas(rows[2])[0].paragraphs[0]
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    espaco = " " * 22
    partes = [
        (f"Convênio nº.: {d.get('numero') or '---'}", not d.get("numero")),
        (f"{espaco}Parcela: {d.get('parcelas') or '---'}", not d.get("parcelas")),
        (f"{espaco}Valor MGI: {fmt_money(d.get('valor_concedente'))}", False),
        (f"{espaco}Valor Contrapartida: {fmt_money(d.get('valor_contrapartida'))}", False),
    ]
    base = deepcopy(p.runs[0]._r) if p.runs else None
    for r in p.runs:
        r._r.getparent().remove(r._r)
    for texto, destacar in partes:
        if base is not None:
            el = deepcopy(base)
            p._p.append(el)
        run = p.runs[-1] if base is not None else p.add_run()
        run.text = texto
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW if destacar else None
    _texto(_celulas_unicas(rows[3])[0].paragraphs[0], f"Objeto: {d.get('objeto') or '---'}", not d.get("objeto"))

    modelo_row = rows[5]._tr
    for extra in rows[6:]:
        extra._tr.getparent().remove(extra._tr)
    ancora = modelo_row
    for i, item in enumerate(itens, start=1):
        novo = deepcopy(modelo_row)
        ancora.addnext(novo)
        ancora = novo
        row = next(r for r in tabela.rows if r._tr is novo)
        cels = _celulas_unicas(row)
        _texto(cels[0].paragraphs[0], f"{i:02d}")
        _preencher_celula_lista(cels[1], item["titulo"], item["situacao"])
        _texto(cels[2].paragraphs[0], f"{i:02d}")
        _preencher_celula_lista(cels[3], "Enviar:", item["providencia"])
    modelo_row.getparent().remove(modelo_row)

    buf = io.BytesIO()
    doc.save(buf)
    numero = re.sub(r"[^\w-]+", "-", d.get("numero") or "sem_numero")
    municipio = (d.get("municipio") or "").upper()
    nome = f"ANEXO - SOLICITAÇÃO DE DOCUMENTOS - CONV {numero}" + (f" - {municipio}" if municipio else "")
    return buf.getvalue(), f"{nome}.docx"


# ─────────────────────────────────────────────────────────────────────────────
# Ofício
# ─────────────────────────────────────────────────────────────────────────────

def _par(doc, partes: List[Tuple[str, bool, bool]], alinhamento=None, espaco_depois: int = 0, recuo: Optional[float] = None):
    p = doc.add_paragraph()
    for texto, negrito, destacar in partes:
        r = p.add_run(texto)
        r.bold = negrito
        r.font.size = Pt(11)
        r.font.name = "Calibri"
        if destacar:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if alinhamento is not None:
        p.alignment = alinhamento
    p.paragraph_format.space_after = Pt(espaco_depois)
    p.paragraph_format.space_before = Pt(0)
    if recuo is not None:
        p.paragraph_format.first_line_indent = Cm(recuo)
    return p


def gerar_oficio(resultado: Dict[str, Any], hoje: Optional[date] = None) -> Tuple[bytes, str]:
    parecer = resultado.get("parecer")
    if not parecer:
        raise _sem_parecer()
    if not parecer.get("solicitacao"):
        raise ValueError("Não há documentos a solicitar: a análise não apontou pendências para diligência.")
    d = parecer["dados"]
    hoje = hoje or date.today()
    numero = d.get("numero") or "[Nº DO CONVÊNIO]"
    municipio = d.get("municipio") or "[MUNICÍPIO]"
    interv_nome = d.get("interveniente_nome") or "[SECRETARIA INTERVENIENTE]"
    interv_sigla = d.get("interveniente_sigla") or "[SIGLA]"
    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    C = WD_ALIGN_PARAGRAPH.CENTER

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)

    _par(doc, [("GOVERNO DO ESTADO DE MINAS GERAIS", True, False)], C)
    _par(doc, [("Minas Gerais Participações S.A", True, False)], C)
    _par(doc, [("Gerência de Convênios", True, False)], C, 12)
    _par(doc, [("Ofício MGI/GECOV nº. ", True, False), (f"[XXX]/{hoje.year}", True, True)])
    _par(doc, [(f"Belo Horizonte, {data_extenso(hoje)}.", False, False)], WD_ALIGN_PARAGRAPH.RIGHT, 12)

    _par(doc, [("Exmo Senhor,", False, False)])
    _par(doc, [("[NOME DO(A) PREFEITO(A)]", False, True)])
    _par(doc, [(f"Prefeito do Município de {municipio}/MG", False, False)])
    _par(doc, [("[ENDEREÇO DA PREFEITURA]", False, True)])
    _par(doc, [("CEP: [XX.XXX-XXX]", False, True), (f" - {municipio}/MG", False, False)], espaco_depois=12)

    _par(doc, [("Assunto: ", True, False),
               (f"Solicitação de documentos complementares - Convênio nº {numero} - Prefeitura Municipal de {municipio}/MG", False, False)], J)
    _par(doc, [("Referência: ", True, False), ("[Caso responda este Ofício, indicar expressamente o Processo nº ", False, False),
               ("XXXX.XX.XXXXXXX/XXXX-XX", False, True), ("].", False, False)], J, 12)

    _par(doc, [("CONVÊNIO: ", True, False), (numero, False, not d.get("numero"))])
    _par(doc, [("CONCEDENTE: ", True, False), ("MGI Minas Gerais Participações S/A", False, False)])
    _par(doc, [("INTERVENIENTE: ", True, False), (interv_nome, False, not d.get("interveniente_nome"))])
    _par(doc, [("CONVENENTE: ", True, False), (f"Prefeitura Municipal de {municipio}/MG", False, False)], espaco_depois=12)

    _par(doc, [(f"Informamos que, em trabalho de análise de prestação de contas final do Convênio nº {numero}, celebrado entre a "
                f"MGI-Minas Gerais Participações S/A, e o Município de {municipio}/MG, com a interveniência da {interv_sigla}, "
                "verificou-se pendência de alguns documentos conforme destacado no “Relatório de Análise de Prestação de "
                "Contas” anexo.", False, False)], J, 8, 1.25)
    _par(doc, [("Gentileza observar na coluna “Providência a ser tomada pelo Convenente” do referido anexo, as medidas que "
                "devem ser adotadas para que sejam sanadas as irregularidades/invalidades destacadas.", False, False)], J, 8, 1.25)
    _par(doc, [("Solicitamos que seja providenciada, no prazo de 15 (quinze) dias corridos após o recebimento desta notificação, "
                "a apresentação dos documentos solicitados ou das justificativas pela impossibilidade de fazê-lo e/ou alegações "
                f"de defesa, {parecer.get('artigo_diligencia')}. A não apresentação dentro do prazo, poderá ensejar pena de "
                "registro da inadimplência no Sistema Integrado da Administração Financeira – SIAFI-MG e reprovação das contas, "
                "com posterior instauração de Tomada de Contas Especial.", False, False)], J, 8, 1.25)
    _par(doc, [("Colocamo-nos à disposição para quaisquer esclarecimentos adicionais que se fizerem necessários por meio do "
                "seguinte contato:", False, False)], J, 0, 1.25)
    _par(doc, [("Nome do Analista: ", False, False), ("[NOME DO ANALISTA]", False, True)])
    _par(doc, [("Telef.: (31) 3915-4885", False, False)])
    _par(doc, [("e-mail: ", False, False), ("[e-mail]@mgipar.com.br", False, True), (" ou convenios@mgipar.com.br", False, False)],
         espaco_depois=18)
    _par(doc, [("Atenciosamente,", False, False)])

    buf = io.BytesIO()
    doc.save(buf)
    rotulo = re.sub(r"[^\w-]+", "-", numero)
    return buf.getvalue(), f"Oficio_diligencia_CV_{rotulo}.docx"
