"""
Saída 2 – Minuta da Análise Documental (AD) nos modelos do Novo TCT da GECOV.

Preenche o modelo Word oficial (Decreto 43.635/2003, Decreto 46.319/2013 ou
Inexecução) com os dados do dossiê. Tudo o que depende de conferência do
analista fica destacado em amarelo; as caixas "Atende / Atende com ressalvas /
Não Atende / Não Consta" só são marcadas quando o Relatório de Validação tem
resultado seguro. Quando há documentos a solicitar (diligência), as seções III
a VI ficam como no modelo, como faz a GECOV.
"""

import io
import re
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from backend.analise_documental.texto import PENDENTE_HUMANO, fmt_money, norm, parse_date

MODELOS_DIR = Path(__file__).parent / "modelos"
MODELOS_AD = {
    "DEC_43635": ("ad_dec43635.docx", "Decreto 43.635/2003"),
    "DEC_46319": ("ad_dec46319.docx", "Decreto 46.319/2013"),
    "INEXECUCAO": ("ad_inexecucao.docx", "Inexecução"),
}

ATRASO_PC = {
    "DEC_43635": "Encaminhamento da Prestação de Contas Final ultrapassando o prazo estabelecido no §5º do art. 26, do Decreto 43.635/2003.",
    "DEC_46319": "Encaminhamento da Prestação de Contas Final ultrapassando o prazo estabelecido no art. 54, §3º do Decreto 46.319/2013.",
}
BASE_INEXECUCAO = {
    "DEC_43635": "art. 12, XIII, alínea a, do Decreto Estadual nº 43.635, de 2003",
    "DEC_46319": "Decreto 46.319/2013 e art. 60, III, da Resolução Conjunta SEGOV/AGE nº 004/2015",
}

TITULO_ATRASO = "Prestação de contas apresentada fora do prazo"
TITULO_DIVERGENCIA_SALDO = "ALERTA DE DIVERGÊNCIA – saldo apurado × saldo bancário"

BLOCO_LICITACAO = ("Processo licitatório",)
BLOCO_OS = ("Ordem de serviço",)
BLOCO_DESPESAS = ("Comprovantes de despesas",)
BLOCO_BANCO = ("Movimentação bancária", "Aplicações/rendimentos", "Contrapartida", "Saldo", "Devoluções",
               "Irregularidades financeiras")
ITENS_EM_BLOCOS = set(BLOCO_LICITACAO + BLOCO_OS + BLOCO_DESPESAS + BLOCO_BANCO + ("Prazo da prestação de contas",))

CAIXA = {"ATENDE": 0, "ATENDE COM RESSALVAS": 1, "NÃO ATENDE": 2, "NÃO CONSTA": 3}
CABECALHO_IRREGULARIDADES = {
    "ATENDE COM RESSALVAS": "Irregularidades ou invalidades:",
    "NÃO ATENDE": "Irregularidades e/ou invalidades:",
    "NÃO CONSTA": "Irregularidades e/ou invalidades:",
}
PREFIXO_CATEGORIA = {
    "a": "Ressalva formal",
    "b": "Pendência documental",
    "c": "Inconsistência financeira (a validar)",
    "d": "A validar",
    "e": "Possível dano ao erário (a validar)",
    "f": "Possível dano ao erário (a validar)",
    "g": "Depende de manifestação técnica",
    "h": "A validar",
    "i": "A validar",
    "j": "Informação insuficiente",
}

Parte = Tuple[str, bool]  # (texto, destacar)


# ─────────────────────────────────────────────────────────────────────────────
# Valor por extenso
# ─────────────────────────────────────────────────────────────────────────────

_UNID = ["", "um", "dois", "três", "quatro", "cinco", "seis", "sete", "oito", "nove", "dez", "onze", "doze",
         "treze", "quatorze", "quinze", "dezesseis", "dezessete", "dezoito", "dezenove"]
_DEZ = ["", "", "vinte", "trinta", "quarenta", "cinquenta", "sessenta", "setenta", "oitenta", "noventa"]
_CEN = ["", "cento", "duzentos", "trezentos", "quatrocentos", "quinhentos", "seiscentos", "setecentos",
        "oitocentos", "novecentos"]
_ESCALAS = [("", ""), ("mil", "mil"), ("milhão", "milhões"), ("bilhão", "bilhões")]


def _ate_mil(n: int) -> str:
    if n == 100:
        return "cem"
    c, r = divmod(n, 100)
    partes = [_CEN[c]] if c else []
    if r:
        if r < 20:
            partes.append(_UNID[r])
        else:
            d, u = divmod(r, 10)
            partes.append(_DEZ[d] + (f" e {_UNID[u]}" if u else ""))
    return " e ".join(partes)


def _inteiro_extenso(n: int) -> str:
    if n == 0:
        return "zero"
    grupos = []
    escala = 0
    while n:
        n, g = divmod(n, 1000)
        if g:
            grupos.append((g, escala))
        escala += 1
    grupos.reverse()
    textos = []
    for g, esc in grupos:
        if esc == 0:
            textos.append(_ate_mil(g))
        elif esc == 1:
            textos.append("mil" if g == 1 else f"{_ate_mil(g)} mil")
        else:
            textos.append(f"{_ate_mil(g)} {_ESCALAS[esc][0] if g == 1 else _ESCALAS[esc][1]}")
    out = textos[0]
    for k in range(1, len(textos)):
        g = grupos[k][0]
        ultimo = k == len(textos) - 1
        out += (" e " if ultimo and (g < 100 or g % 100 == 0) else ", ") + textos[k]
    return out


def valor_por_extenso(valor: float) -> str:
    centavos_total = int(round(abs(valor) * 100))
    inteiro, centavos = divmod(centavos_total, 100)
    partes = []
    if inteiro:
        moeda = "real" if inteiro == 1 else ("de reais" if inteiro % 1_000_000 == 0 else "reais")
        partes.append(f"{_inteiro_extenso(inteiro)} {moeda}")
    if centavos:
        partes.append(f"{_inteiro_extenso(centavos)} {'centavo' if centavos == 1 else 'centavos'}")
    return " e ".join(partes) if partes else "zero real"


# ─────────────────────────────────────────────────────────────────────────────
# Coleta dos dados do dossiê (compartilhada com a conciliação em Excel)
# ─────────────────────────────────────────────────────────────────────────────

def _valor(campos: Optional[Dict], nome: str):
    return ((campos or {}).get(nome) or {}).get("valor")


def _confirmado(campos: Optional[Dict], nome: str) -> bool:
    return ((campos or {}).get(nome) or {}).get("status") == "confirmado"


def fls(localizacao: Optional[str]) -> Tuple[str, bool]:
    """Folha física quando o carimbo foi lido; senão arquivo/página (a conferir)."""
    if not localizacao:
        return "", False
    m = re.search(r"\bfls?\.\s*([\d][\d\s,–-]*)", localizacao)
    if m:
        return m.group(1).strip(" ,"), True
    arq = re.sub(r"\.pdf\b", "", localizacao.split(",")[0], flags=re.IGNORECASE)
    pag = re.search(r"pp?\.\s*([\d–-]+)", localizacao)
    return (f"{arq}, p. {pag.group(1)}" if pag else arq), False


def _fls_intervalo(docs: List[Dict]) -> Tuple[str, bool]:
    if not docs:
        return "", False
    ini, ok1 = fls(docs[0]["localizacao"])
    if len(docs) == 1:
        return ini, ok1
    fim, ok2 = fls(docs[-1]["localizacao"])
    return f"{ini} a {fim}", ok1 and ok2


def coletar_dados(resultado: Dict[str, Any]) -> Dict[str, Any]:
    dados = resultado.get("dados_convenio") or {}
    cruz = resultado.get("cruzamentos") or {}
    rel = resultado.get("relatorio_validacao") or {}
    ctrl = cruz.get("controle") or {}
    ativos = [d for d in resultado.get("inventario") or [] if not d.get("duplicado_de")]
    por_tipo: Dict[str, List[Dict]] = {}
    for d in ativos:
        por_tipo.setdefault(d["tipo"], []).append(d)

    regime = rel.get("regime") or {}
    itens = {i["item"]: i for i in rel.get("itens") or []}
    ocorrencias = rel.get("ocorrencias") or []

    pact_conc = _valor(dados, "valor_concedente")
    pact_cp = _valor(dados, "valor_contrapartida") or 0.0
    base_pt = (pact_conc or 0.0) + pact_cp
    perc_conc = (pact_conc / base_pt) if pact_conc and base_pt else None

    repasses = ctrl.get("repasses")
    contrapartida = ctrl.get("contrapartida") or 0.0
    rendimentos = ctrl.get("rendimentos_extratos")
    if rendimentos is None:
        rendimentos = (cruz.get("conhecimentos_receita") or {}).get("total_rendimentos") or 0.0
    despesas = ctrl.get("despesas_notas_fiscais") or 0.0
    devolucoes = cruz.get("devolucoes") or []
    total_devolvido = cruz.get("total_devolvido") or sum(d.get("valor") or 0 for d in devolucoes)

    df: Dict[str, Optional[float]] = {}
    if repasses is not None:
        total_rec = repasses + contrapartida + rendimentos
        saldo = total_rec - despesas
        prop = saldo * perc_conc if perc_conc is not None else None
        parcial = prop - total_devolvido if prop is not None else None
        aporte_menor = max(0.0, pact_cp - contrapartida) * perc_conc if perc_conc is not None else 0.0
        df = {
            "concedente": repasses, "contrapartida": contrapartida, "rendimentos": rendimentos,
            "recursos_proprios": 0.0, "outras_receitas": 0.0, "total_receitas": total_rec,
            "cheques": 0.0, "teds": despesas, "total_despesas": despesas, "saldo": saldo,
            "saldo_proporcional": prop, "devolucao_efetuada": total_devolvido, "parcial": parcial,
            "aporte_menor": aporte_menor,
            "final": (parcial + aporte_menor) if parcial is not None else None,
        }

    notas_sem_valor = [n for n in cruz.get("notas_fiscais") or [] if n.get("valor_bruto") is None]
    df_incerto = bool(notas_sem_valor) or any(o["titulo"] == TITULO_DIVERGENCIA_SALDO for o in ocorrencias) \
        or not df or df.get("final") is None

    parecer = (por_tipo.get("PARECER_TECNICO") or [None])[0]
    oficio = (por_tipo.get("OFICIO") or [None])[0]
    lic = (por_tipo.get("LICITACAO") or [None])[0]
    contratos = [c for c in por_tipo.get("CONTRATO") or [] if _valor(c.get("extracao"), "contratada")] \
        or por_tipo.get("CONTRATO") or []

    extracao_por_doc = {d["id"]: d.get("extracao") or {} for d in ativos}
    notas = []
    for n in cruz.get("notas_fiscais") or []:
        e = extracao_por_doc.get(n.get("documento"), {})
        notas.append({
            **n,
            "inss": _valor(e, "retencao_inss"),
            "ir": _valor(e, "retencao_ir"),
            "iss": _valor(e, "retencao_iss"),
        })

    return {
        "dados": dados,
        "cruz": cruz,
        "rel": rel,
        "regime": regime,
        "regime_codigo": regime.get("codigo"),
        "itens": itens,
        "ocorrencias": ocorrencias,
        "por_tipo": por_tipo,
        "perc_conc": perc_conc,
        "pact_conc": pact_conc,
        "pact_cp": pact_cp,
        "df": df,
        "df_incerto": df_incerto,
        "devolucoes": devolucoes,
        "total_devolvido": total_devolvido,
        "tarifas": (cruz.get("evidencias_bancarias") or {}).get("tarifas") or [],
        "parecer": parecer,
        "oficio": oficio,
        "licitacao": lic,
        "contrato": contratos[0] if contratos else None,
        "notas": notas,
        "atrasado": next((o for o in ocorrencias if o["titulo"] == TITULO_ATRASO), None),
    }


def sugerir_modelo(resultado: Dict[str, Any]) -> Dict[str, Any]:
    info = coletar_dados(resultado)
    codigo = info["regime_codigo"]
    if codigo in ("DEC_43635", "DEC_46319"):
        modelo, motivo = codigo, f"Regime identificado: {info['regime'].get('nome')}."
    else:
        modelo = "DEC_46319"
        motivo = ("Regime jurídico não identificado com segurança ou sem modelo TCT próprio "
                  f"({info['regime'].get('nome') or 'não identificado'}); confirme o modelo antes de usar.")
    parecer_txt = " ".join(str(_valor((info["parecer"] or {}).get("extracao"), k) or "")
                           for k in ("conclusao", "percentual_execucao", "glosas_ou_irregularidades"))
    sinais = []
    if re.search(r"NAO\s+(FOI\s+)?EXECUTAD|INEXECU|(?<![\d,.])0\s*%", norm(parecer_txt)):
        sinais.append("o Parecer Técnico indica não execução")
    if not info["notas"] and not (info["cruz"].get("evidencias_bancarias") or {}).get("pagamentos"):
        sinais.append("não há notas fiscais nem pagamentos identificados")
    return {
        "modelo_sugerido": modelo,
        "motivo": motivo,
        "inexecucao_sugerida": bool(sinais),
        "motivo_inexecucao": "; ".join(sinais),
        "modelos": {k: v[1] for k, v in MODELOS_AD.items()},
    }


# ─────────────────────────────────────────────────────────────────────────────
# Manipulação do documento
# ─────────────────────────────────────────────────────────────────────────────

def _cor(run: Run, destacar: Optional[bool]):
    if destacar is not None:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW if destacar else None


def _set_runs(p: Paragraph, partes: List[Tuple], bold_rotulo: bool = False):
    """Reescreve o parágrafo preservando a formatação do primeiro run. partes: (texto, destacar[, negrito])."""
    runs = p.runs
    base = deepcopy(runs[0]._r) if runs else None
    for r in runs:
        r._r.getparent().remove(r._r)
    for parte in partes:
        texto, destacar = parte[0], parte[1]
        if base is not None:
            el = deepcopy(base)
            p._p.append(el)
            run = Run(el, p)
            run.text = texto
        else:
            run = p.add_run(texto)
        _cor(run, destacar)
        if len(parte) > 2 and parte[2] is not None:
            run.bold = parte[2]


def _set_text(p: Paragraph, texto: str, destacar: Optional[bool] = None):
    runs = p.runs
    if runs:
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        runs[0].text = texto
        _cor(runs[0], destacar)
    else:
        _cor(p.add_run(texto), destacar)


def _remover(p: Paragraph):
    el = p._p
    if el.getparent() is not None:
        el.getparent().remove(el)


def _novo_paragrafo_apos(ref: Paragraph, partes: List[Tuple], modelo: Optional[Paragraph] = None) -> Paragraph:
    novo = deepcopy((modelo or ref)._p)
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for el in novo.findall(".//" + qn(tag)):
            el.getparent().remove(el)
    ppr = novo.find(qn("w:pPr"))
    if ppr is not None:
        for jc in ppr.findall(qn("w:jc")):
            ppr.remove(jc)
    ref._p.addnext(novo)
    par = Paragraph(novo, ref._parent)
    _set_runs(par, partes)
    return par


def _lista_apos(ref: Paragraph, linhas: List[Parte], modelo: Optional[Paragraph] = None) -> Paragraph:
    ultimo = ref
    for texto, destacar in linhas:
        ultimo = _novo_paragrafo_apos(ultimo, [(texto, destacar, False)], modelo)
    return ultimo


def _tem_caixa(p: Paragraph) -> bool:
    return p._p.find(".//" + qn("w:checkBox")) is not None


def _marcar(p: Paragraph, indice: Optional[int]):
    caixas = p._p.findall(".//" + qn("w:checkBox"))
    for cb in caixas:
        for old in cb.findall(qn("w:checked")):
            cb.remove(old)
    if indice is not None and indice < len(caixas):
        caixas[indice].append(OxmlElement("w:checked"))


def _celulas(doc) -> List[_Cell]:
    vistos, out = set(), []
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                if id(c._tc) not in vistos:
                    vistos.add(id(c._tc))
                    out.append(c)
    return out


def _achar(doc, padrao: str) -> Tuple[Optional[_Cell], Optional[Paragraph]]:
    rx = re.compile(padrao)
    for c in _celulas(doc):
        for p in c.paragraphs:
            if rx.search(norm(p.text)):
                return c, p
    return None, None


def _depois(cell: _Cell, ref: Paragraph, padrao: Optional[str] = None, caixa: bool = False,
            ate: Optional[str] = None) -> Optional[Paragraph]:
    paras = cell.paragraphs
    idx = next((i for i, q in enumerate(paras) if q._p is ref._p), None)
    if idx is None:
        return None
    rx = re.compile(padrao) if padrao else None
    rx_ate = re.compile(ate) if ate else None
    for q in paras[idx + 1:]:
        t = norm(q.text)
        if rx_ate and rx_ate.search(t):
            return None
        if caixa and _tem_caixa(q):
            return q
        if rx and rx.search(t):
            return q
    return None


def _sub_texto(p: Paragraph, padrao: str, novo: str, destacar: Optional[bool] = None, flags=re.IGNORECASE) -> bool:
    texto = p.text
    alterado = re.sub(padrao, novo, texto, flags=flags)
    if alterado != texto:
        _set_text(p, alterado, destacar)
        return True
    return False


# ─────────────────────────────────────────────────────────────────────────────
# Conteúdo
# ─────────────────────────────────────────────────────────────────────────────

def _sem_pendente(texto: str) -> str:
    return texto.replace(PENDENTE_HUMANO, "").strip()


def _linhas_ocorrencias(ocorrencias: List[Dict], compacto: bool = False) -> List[Parte]:
    grupos: Dict[Tuple[str, str], List[Dict]] = {}
    for o in ocorrencias:
        grupos.setdefault((o["categoria"], o["titulo"]), []).append(o)
    linhas = []
    for (cat, titulo), lista in grupos.items():
        o = lista[0]
        if compacto:
            sufixo = f" ({len(lista)} ocorrências)" if len(lista) > 1 else f": {_sem_pendente(o['descricao'])}"
            linhas.append((f"{titulo}{sufixo}", cat != "a"))
            continue
        prefixo = PREFIXO_CATEGORIA.get(cat, "A validar")
        if len(lista) > 1:
            texto = f"{prefixo} – {titulo} ({len(lista)} ocorrências; ex.: {o['descricao']})"
        else:
            texto = f"{prefixo} – {titulo}: {o['descricao']}"
        if o.get("valor") is not None and len(lista) == 1 and "R$" not in o["descricao"]:
            texto += f" Valor: {fmt_money(o['valor'])}."
        if o.get("fundamento"):
            texto += f" Fundamento: {o['fundamento']}."
        fontes = [f for f in (o.get("fontes") or []) if f][:2]
        if fontes:
            texto += f" (Ref.: {'; '.join(fontes)})"
        linhas.append((texto, cat != "a"))
    return linhas


def _resultado_bloco(info: Dict, itens: Tuple[str, ...]) -> Tuple[Optional[int], List[Dict]]:
    achados = [info["itens"][i] for i in itens if i in info["itens"]]
    if not achados:
        return None, []
    principal = achados[0]
    if principal["resultado"] == "NÃO CONSTA":
        return CAIXA["NÃO CONSTA"], achados
    if any(a["resultado"] == "ATENDE COM RESSALVAS" for a in achados):
        return CAIXA["ATENDE COM RESSALVAS"], achados
    if all(a["resultado"] == "ATENDE" for a in achados):
        return CAIXA["ATENDE"], achados
    return None, achados


def _texto_campo(info: Dict, nome: str, rotulo: str, formatar=None) -> List[Tuple]:
    dados = info["dados"]
    v = _valor(dados, nome)
    if v is None or v == "":
        return [(rotulo, False, True), (" NÃO IDENTIFICADO – conferir", True, False)]
    texto = formatar(v) if formatar else str(v)
    return [(rotulo, False, True), (f" {texto}", not _confirmado(dados, nome), False)]


def _pct(v: float) -> str:
    return f"{v:.2f}".replace(".", ",") + "%"


def _municipio(v: str) -> str:
    return re.sub(r"^(munic[ií]pio|prefeitura municipal)\s+de\s+", "", str(v), flags=re.IGNORECASE).strip()


def _preencher_dados(doc, info: Dict, modelo: str):
    dados = info["dados"]
    campos = [
        (r"^NUMERO DO CONVENIO", "numero_convenio", "Número do Convênio:", None),
        (r"^MUNICIPIO", "convenente", "Município:", _municipio),
        (r"^CONVENENTE:", "convenente", "Convenente:", None),
        (r"^OBJETO:", "objeto", "Objeto:", None),
        (r"^VALOR CONCEDENTE", "valor_concedente", "Valor Concedente:", fmt_money),
        (r"^VALOR CONTRAPARTIDA", "valor_contrapartida", "Valor Contrapartida:", fmt_money),
        (r"^VALOR TOTAL PACTUADO", "valor_total", "Valor Total Pactuado:", fmt_money),
    ]
    for padrao, nome, rotulo, fmt in campos:
        _, p = _achar(doc, padrao)
        if p is not None:
            _set_runs(p, _texto_campo(info, nome, rotulo, fmt))

    _, p = _achar(doc, r"^SECRETARIA INTERVENIENTE")
    if p is not None:
        interv = _valor(dados, "interveniente")
        if interv:
            _set_runs(p, [("Secretaria Interveniente:", False, True), (f" {interv}", not _confirmado(dados, "interveniente"), False)])
        else:
            conc = _valor(dados, "concedente")
            _set_runs(p, [("Secretaria Interveniente:", False, True),
                          (f" {conc + ' (concedente original – ' if conc else ' NÃO IDENTIFICADA ('}conferir a interveniente do TCT)", True, False)])

    _, p = _achar(doc, r"^VIGENCIA:")
    if p is not None:
        ini, fim = _valor(dados, "vigencia_inicio"), _valor(dados, "vigencia_fim")
        ok = _confirmado(dados, "vigencia_inicio") and _confirmado(dados, "vigencia_fim")
        _set_runs(p, [("Vigência:", False, True), (f" {ini or '---'} a {fim or '---'}", not (ok and ini and fim), False)])

    cell, p = _achar(doc, r"^PRESTACAO DE CONTAS FINAL")
    if p is not None:
        e = (info["oficio"] or {}).get("extracao") or {}
        num, data = _valor(e, "numero"), _valor(e, "data")
        prot_num, prot_data = _valor(e, "numero_protocolo"), _valor(e, "data_protocolo")
        d_of, d_pr = parse_date(data), parse_date(prot_data)
        siged = prot_num or "---"
        if d_pr and (not d_of or d_pr >= d_of):
            siged += f" de {prot_data}"
        _set_runs(p, [
            ("Prestação de Contas Final: OFÍCIO Nº ", False, True),
            (f"{num or '---'} DE {data or '---/---/----'}", not (num and data), True),
            (" SIGED ", False, True),
            (siged, True, True),
        ])
        caixa = _depois(cell, p, caixa=True)
        if caixa is not None:
            item = info["itens"].get("Prazo da prestação de contas") or {}
            _marcar(caixa, CAIXA.get(item.get("resultado")))
        _preencher_ressalva_prazo(cell, info, modelo)


def _preencher_ressalva_prazo(cell: _Cell, info: Dict, modelo: str):
    cab = next((q for q in cell.paragraphs if norm(q.text).startswith("RESSALVAS")), None)
    if cab is None:
        return
    variantes = [q for q in cell.paragraphs if norm(q.text).startswith("ENCAMINHAMENTO DA PRESTACAO")
                 or norm(q.text) == "OU"]
    regime = modelo if modelo in ATRASO_PC else info["regime_codigo"]
    atraso = info["atrasado"]
    item = info["itens"].get("Prazo da prestação de contas") or {}
    if atraso:
        _set_runs(cab, [("Ressalvas:", False, True)])
        texto = ATRASO_PC.get(regime) or " OU ".join(ATRASO_PC.values())
        _lista_apos(cab, [(f"{texto} {atraso['descricao']}", "ofício" in atraso["descricao"] or regime not in ATRASO_PC)],
                    modelo=variantes[0] if variantes else None)
    elif item.get("resultado") == "ATENDE":
        _set_runs(cab, [("Ressalvas: ", False, True), ("não identificadas quanto ao prazo de apresentação.", False, False)])
    else:
        _set_runs(cab, [("Ressalvas: ", False, True),
                        (f"prazo pendente de validação – {item.get('observacao') or 'dados insuficientes'}", True, False)])
    for q in variantes:
        _remover(q)


def _preencher_bloco(doc, padrao: str, cabecalho: Optional[List[Tuple]], info: Dict, itens: Tuple[str, ...]):
    cell, p = _achar(doc, padrao)
    if p is None:
        return
    if cabecalho:
        _set_runs(p, cabecalho)
    caixa = _depois(cell, p, caixa=True)
    indice, achados = _resultado_bloco(info, itens)
    if caixa is not None:
        _marcar(caixa, indice)
    ocorr = [o for o in info["ocorrencias"] if o["item"] in itens]
    linhas = _linhas_ocorrencias(ocorr)
    for a in achados:
        if a["resultado"] not in CAIXA and a["resultado"] != "REGIME IDENTIFICADO":
            linhas.insert(0, (f"Pendente de validação humana ({a['item']}): {a['observacao']}", True))
    if not linhas:
        linhas = [("Não foram identificadas irregularidades nos documentos analisados.", False)]

    cab = _depois(cell, p, r"^IRREGULARIDADES")
    lista = _depois(cell, cab, r"^LISTAR AS IRREGULARIDADES") if cab is not None else None
    if cab is not None:
        _set_runs(cab, [("Irregularidades e/ou invalidades (ressalvas e/ou danos):", False, True)])
    if lista is not None:
        _set_runs(lista, [(linhas[0][0], linhas[0][1], False)])
        _lista_apos(lista, linhas[1:])
    elif caixa is not None and indice != CAIXA["ATENDE"]:
        _lista_apos(caixa, linhas)


def _conta_bancaria(conta: str) -> Tuple[str, str, str]:
    banco = re.search(r"banco\D{0,4}(\d{3})", conta, re.IGNORECASE)
    ag = re.search(r"ag(?:[êe]ncia|\.)?\s*:?\s*([\dxX][\dxX.\-]*)", conta, re.IGNORECASE)
    cc = re.search(r"(?:conta|c/c)\s*(?:corrente)?\s*:?\s*([\dxX][\dxX.\-]*)", conta, re.IGNORECASE)
    if banco and ag and cc:
        return banco.group(1), ag.group(1), cc.group(1)
    partes = conta.split()
    if len(partes) == 3:
        return partes[0], partes[1], partes[2]
    return "---", "---", conta or "---"


def _preencher_documentacao(doc, info: Dict):
    lic = (info["licitacao"] or {}).get("extracao") or {}
    ctr = (info["contrato"] or {}).get("extracao") or {}
    docs_lic = ([info["licitacao"]] if info["licitacao"] else []) + ([info["contrato"]] if info["contrato"] else [])
    fl, fl_ok = _fls_intervalo(docs_lic)
    processo = _valor(lic, "numero_processo") or _valor(lic, "modalidade")
    numero_ctr = _valor(ctr, "numero")
    empresa = _valor(ctr, "contratada") if _confirmado(ctr, "contratada") else (_valor(lic, "vencedor") or _valor(ctr, "contratada"))
    valor_ctr = _valor(ctr, "valor") or _valor(lic, "valor_adjudicado")
    vig_ctr = _valor(ctr, "prazo_execucao")
    modalidade = _valor(lic, "modalidade")
    _preencher_bloco(doc, r"^PROCESSO LICITATORIO N", [
        ("Processo Licitatório nº ", False, True), (f"{processo or '---'}", not processo, True),
        (f" ({modalidade})" if modalidade and modalidade != processo else "", False, True),
        (" - Contrato nº ", False, True), (numero_ctr or "---", not numero_ctr, True),
        (" - Empresa: ", False, True), (empresa or "---", not empresa, True),
        (" - ", False, True), (fmt_money(valor_ctr) if valor_ctr else "R$ ---", not valor_ctr, True),
        (" - Vigência ", False, True), (vig_ctr or "---", not vig_ctr, True),
        (" (Fls.: ", False, True), (fl or "---", not fl_ok, True), (")", False, True),
    ], info, BLOCO_LICITACAO)

    os_docs = info["por_tipo"].get("ORDEM_SERVICO") or []
    fl, fl_ok = _fls_intervalo(os_docs[:1])
    _preencher_bloco(doc, r"^ORDEM DE SERVICO \(FLS", [
        ("Ordem de Serviço (Fls.: ", False, True), (fl or "---", not fl_ok, True), (")", False, True),
    ], info, BLOCO_OS)

    nf_docs = info["por_tipo"].get("NOTA_FISCAL") or []
    fl, fl_ok = _fls_intervalo(nf_docs)
    _preencher_bloco(doc, r"^COMPROVANTES DE DESPESAS \(FLS", [
        ("Comprovantes de Despesas (Fls.: ", False, True), (fl or "---", not fl_ok, True), (")", False, True),
    ], info, BLOCO_DESPESAS)

    banco, ag, cc = _conta_bancaria(str(_valor(info["dados"], "banco_agencia_conta") or ""))
    ext = info["por_tipo"].get("EXTRATOS") or []
    fl, fl_ok = _fls_intervalo(ext)
    conta_ok = _confirmado(info["dados"], "banco_agencia_conta")
    _preencher_bloco(doc, r"^MOVIMENTACAO BANCARIA", [
        ("Movimentação Bancária Banco: ", False, True), (banco, not conta_ok, True),
        (" Ag: ", False, True), (ag, not conta_ok, True), (" C/C ", False, True), (cc, not conta_ok, True),
        (" (Fls.: ", False, True), (fl or "---", not fl_ok, True), (")", False, True),
    ], info, BLOCO_BANCO)

    # Inexecução: blocos sem cabeçalho de dados
    for padrao, itens in ((r"^PROCESSO LICITATORIO \(", BLOCO_LICITACAO), (r"^ORDEM DE SERVICO \( OBJETO", BLOCO_OS),
                          (r"^COMPROVANTES DE DESPESAS \( OBJETO", BLOCO_DESPESAS)):
        _preencher_bloco(doc, padrao, None, info, itens)


def _consideracoes(info: Dict) -> List[Parte]:
    linhas: List[Parte] = []
    reg = info["regime"]
    if reg:
        item = info["itens"].get("Regime jurídico") or {}
        linhas.append((f"Regime jurídico aplicável: {reg.get('nome')}. {reg.get('motivo') or ''}".strip(),
                       item.get("resultado") != "REGIME IDENTIFICADO"))
    faltantes = [c for c in info["rel"].get("checklist_documental") or [] if c.get("situacao") == "NÃO CONSTA"]
    for c in faltantes:
        fund = f" ({c['fundamento']})" if c.get("fundamento") else ""
        linhas.append((f"Documento não localizado no processo: {c['documento']}{fund}.", True))
    outras = [o for o in info["ocorrencias"] if o["item"] not in ITENS_EM_BLOCOS and o["item"] != "Inventário documental"]
    linhas.extend(_linhas_ocorrencias(outras))
    if info["df_incerto"]:
        ctrl = info["cruz"].get("controle") or {}
        linhas.append((
            "Os valores do Demonstrativo Financeiro foram apurados pelo sistema a partir das notas fiscais, repasses, "
            f"contrapartida e rendimentos lidos (saldo apurado {fmt_money(ctrl.get('saldo_apurado'))}; último saldo bancário "
            f"{fmt_money(ctrl.get('saldo_final_extrato_aplicacao'))}) e dependem da conciliação financeira da GECOV.", True))
    return linhas


def _preencher_consideracoes(doc, info: Dict, numero: Optional[str]):
    linhas = _consideracoes(info)
    _, cab = _achar(doc, r"^CONSIDERACOES GERAIS")
    if cab is not None:
        _set_runs(cab, [("CONSIDERAÇÕES GERAIS:", False, True)])
    _, instr = _achar(doc, r"^QUANDO O ANALISTA CONSIDERAR")
    if instr is not None:
        _remover(instr)
    _, base = _achar(doc, r"^COM BASE N[OA]")
    if base is None:
        return
    if "XXX/XXXX" in base.text.upper() or "FLS. (XXX)" in base.text.upper():
        oficio_fl, _ = fls((info["oficio"] or {}).get("localizacao"))
        texto = re.sub(r"Fls\. \(xxx\)", f"Fls. ({oficio_fl or 'xxx'})", base.text, flags=re.IGNORECASE)
        texto = re.sub(r"Convênio nº xxx/xxxx", f"Convênio nº {numero or 'xxx/xxxx'}", texto, flags=re.IGNORECASE)
        _set_text(base, texto)
    if linhas:
        _lista_apos(base, linhas)


def _preencher_parecer(doc, info: Dict):
    parecer = info["parecer"]
    e = (parecer or {}).get("extracao") or {}
    num, data = _valor(e, "numero"), _valor(e, "data")
    conclusao, glosas = _valor(e, "conclusao"), _valor(e, "glosas_ou_irregularidades")
    fl, fl_ok = fls((parecer or {}).get("localizacao"))

    cell, p = _achar(doc, r"^PARECER TECNICO N")
    if p is not None:
        _set_runs(p, [("Parecer Técnico nº ", False, True), (num or "xxxx", True, True),
                      (" (Fls.: ", False, True), (fl or "---", not fl_ok, True), (")", False, True)])
        texto = _depois(cell, p, r"^O PARECER TECNICO DE N")
        if texto is not None:
            if not parecer:
                _set_text(texto, "Parecer Técnico não localizado nos documentos analisados – a conclusão sobre a execução "
                                 "física depende de manifestação técnica.", True)
            elif conclusao:
                _set_text(texto, f"O Parecer Técnico de nº {num or 'xxxxxx'}, de {data or 'xx/xx/xxxx'}, de fls. {fl or 'xxxxxx'}, "
                                 "que analisou a execução da obra conveniada, conforme documentação apresentada pelo convenente "
                                 f"e Relatório de Monitoramento/Vistoria “in loco”, concluiu que, no aspecto técnico, “{conclusao}”.", True)
            else:
                t = texto.text
                t = re.sub(r"de xx/xx/xxxx", f"de {data or 'xx/xx/xxxx'}", t)
                t = re.sub(r"fls\. Xxxxxxx", f"fls. {fl or 'Xxxxxxx'}", t)
                t += " (CONCLUSÃO DO PARECER NÃO LIDA AUTOMATICAMENTE – TRANSCREVER.)"
                _set_text(texto, t, True)
        apontar = _depois(cell, p, r"^APONTAR ALGUM PONTO")
        if apontar is not None:
            if glosas:
                _set_text(apontar, f"Pontos apontados no Parecer Técnico: {glosas}", True)
            else:
                _remover(apontar)

    cell, p = _achar(doc, r"^RELATORIO DE MONITORAMENTO E VISTORIA \(FLS")
    if p is not None:
        _set_runs(p, [("Relatório de Monitoramento e Vistoria (Fls.: ", False, True), (fl or "---", not fl_ok, True),
                      (")", False, True)])


def _preencher_gecon(doc, info: Dict):
    df = info["df"]
    e = (info["parecer"] or {}).get("extracao") or {}
    _, nt = _achar(doc, r"^NOTA TECNICA N")
    if nt is not None:
        glosado = _valor(e, "valor_glosado")
        if glosado:
            _sub_texto(nt, r"R\$\s?X+", fmt_money(glosado), True)
        else:
            _remover(nt)

    def anexar(padrao: str, texto: str, destacar: bool):
        _, p = _achar(doc, padrao)
        if p is not None:
            rotulo = p.text.strip()
            _set_runs(p, [(rotulo, False), (f" {texto}", destacar, False)])

    anexar(r"^SALDO DA CONTA DE CONVENIO", fmt_money(df.get("saldo")) if df else "a apurar.", True)
    anexar(r"^RENDIMENTOS DO VALOR NAO APLICADO", "a apurar na conciliação financeira (calculadora do cidadão).", True)
    tarifas = sum(t.get("valor") or 0 for t in info["tarifas"])
    anexar(r"^TARIFAS BANCARIAS INDEVIDAS",
           fmt_money(tarifas) + " (debitadas na conta do convênio)." if tarifas else "não identificadas nos extratos lidos.", True)
    if info["perc_conc"] is not None:
        pc = info["perc_conc"] * 100
        anexar(r"^PROPORCIONALIDADE:", f"concedente {_pct(pc)} / convenente {_pct(100 - pc)} (valores pactuados: "
                                       f"{fmt_money(info['pact_conc'])} e {fmt_money(info['pact_cp'])}).", False)
    if info["devolucoes"]:
        devs = "; ".join(f"{fmt_money(d.get('valor'))} em {d.get('data') or '—'}" for d in info["devolucoes"])
        anexar(r"^DEVOLUCOES EFETUADAS", f"{devs} (total {fmt_money(info['total_devolvido'])}).", False)
    else:
        anexar(r"^DEVOLUCOES EFETUADAS", "nenhuma devolução localizada.", True)

    cell, alt1 = _achar(doc, r"^SENDO ASSIM, CONSIDERANDO QUE OS SALDOS")
    _, alt2 = _achar(doc, r"^SENDO ASSIM, VERIFICAMOS SALDO HISTORICO")
    if alt1 is None or alt2 is None:
        return
    ou = _depois(cell, alt1, r"^OU$")
    final = (df or {}).get("final")
    valor_txt = f"{fmt_money(abs(final))} ({valor_por_extenso(final)})" if final is not None else "R$xxxxxxxxx (xxxxxxxx)"
    if final is not None and final <= 0:
        mantido, removido = alt1, alt2
    else:
        mantido, removido = alt2, alt1
    t = re.sub(r"\s*\(ESSA PARTE DEVE SER COLOCADA[^)]*\)", "", mantido.text)
    t = re.sub(r"R\$\s?x+\s*\(x+\)", valor_txt, t, flags=re.IGNORECASE)
    _set_text(mantido, t, True)
    _remover(removido)
    if ou is not None:
        _remover(ou)


def _linha_tabela(doc, padrao: str):
    rx = re.compile(padrao)
    for t in doc.tables:
        for row in t.rows:
            if rx.search(norm(row.cells[0].text)):
                return t, row
    return None, None


def _valor_celula(row, texto: str, destacar: bool):
    cell = row.cells[-1]
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        _remover(extra)
    _set_text(p, texto, destacar)


def _preencher_df(doc, info: Dict):
    df = info["df"]
    incerto = info["df_incerto"]
    if not df:
        return
    linhas = [
        (r"^CONCEDENTE$", "concedente", False),
        (r"^CONTRAPARTIDA$", "contrapartida", False),
        (r"^RENDIMENTOS$", "rendimentos", True),
        (r"^RECURSOS PROPRIOS", "recursos_proprios", False),
        (r"^OUTRAS RECEITAS", "outras_receitas", False),
        (r"^TOTAL DE RECEITAS", "total_receitas", incerto),
        (r"^CHEQUES", "cheques", True),
        (r"^TED", "teds", incerto),
        (r"^TOTAL DESPESAS", "total_despesas", incerto),
        (r"^\(=\) SALDO", "saldo", incerto),
        (r"^\(\*\) CALCULO SALDO", "saldo_proporcional", incerto),
        (r"^\(-\) DEVOLUCAO EFETUADA", "devolucao_efetuada", False),
        (r"^\(\+\) APORTE DE CONTRAPARTIDA A MENOR", "aporte_menor", False),
    ]
    for padrao, chave, destacar in linhas:
        _, row = _linha_tabela(doc, padrao)
        if row is not None and df.get(chave) is not None:
            _valor_celula(row, fmt_money(df[chave]), destacar)
    for padrao in (r"^\(\+\) REND\. NAO AUFERIDO", r"^\(\+\) NOTA TECNICA", r"^\(-\) OUTRAS DEVOLUCOES", r"^\(\+/-\) OUTROS",
                   r"^\(-\) DEVOLUCAO OFICIO"):
        for t in doc.tables:
            for row in t.rows:
                if re.search(padrao, norm(row.cells[0].text)):
                    _valor_celula(row, "R$ 0,00 (a apurar)", True)

    for padrao, chave in ((r"^\(=\) RESTITUIR/COBRANCA AO CONVENENTE \(PARCIAL\)", "parcial"),
                          (r"^\(=\) RESTITUIR/COBRANCA AO CONVENENTE$", "final")):
        _, row = _linha_tabela(doc, padrao)
        v = df.get(chave)
        if row is None or v is None:
            continue
        rotulo = "(=) Cobrança ao Convenente" if v > 0 else "(=) Restituir ao Convenente"
        if chave == "parcial":
            rotulo += " (Parcial)"
        _set_text(row.cells[0].paragraphs[0], rotulo, True)
        _valor_celula(row, fmt_money(abs(v)), True)

    # Inexecução
    _, row = _linha_tabela(doc, r"^REPASSE MGI")
    if row is not None:
        _valor_celula(row, fmt_money(df["concedente"]), False)
    _, row = _linha_tabela(doc, r"^\(-\) DEVOLUCAO A MGI")
    if row is not None:
        datas = [d.get("data") for d in info["devolucoes"] if d.get("data")]
        if datas:
            _sub_texto(row.cells[0].paragraphs[0], r"xx/xx/xx", datas[-1])
        _valor_celula(row, fmt_money(info["total_devolvido"]), False)
    for padrao in (r"^\(\+\) RENDIMENTO SELIC", r"^\(=\) VALOR A DEVOLVER", r"^SALDO EM", r"^\(\+\) ATUALIZACAO DE SALDO",
                   r"^\(=\) COBRANCA VALOR PRESENTE"):
        _, row = _linha_tabela(doc, padrao)
        if row is not None:
            _valor_celula(row, "R$ (calcular – SELIC)", True)


def _preencher_devolucoes(doc, info: Dict):
    tabela, cab = _linha_tabela(doc, r"^DATA$|^REFERENCIA$")
    if tabela is None:
        return
    linhas = list(tabela.rows)
    i = next(k for k, r in enumerate(linhas) if r._tr is cab._tr)
    if i + 2 >= len(linhas):
        return
    modelo, total = linhas[i + 1], linhas[i + 2]
    if not norm(total.cells[0].text).startswith("TOTAL"):
        return
    for d in info["devolucoes"]:
        novo = deepcopy(modelo._tr)
        total._tr.addprevious(novo)
        row = [r for r in tabela.rows if r._tr is novo][0]
        unicas = []
        for c in row.cells:
            if not unicas or c._tc is not unicas[-1]._tc:
                unicas.append(c)
        fl, fl_ok = fls(d.get("localizacao"))
        _set_text(unicas[0].paragraphs[0], f"{d.get('data') or '—'} / fls. {fl}", not fl_ok)
        if len(unicas) > 2:
            _set_text(unicas[1].paragraphs[0], f"DAE – {d.get('natureza') or 'natureza não identificada'}", False)
        _set_text(unicas[-1].paragraphs[0], fmt_money(d.get("valor")), not d.get("comprovada"))
    if info["devolucoes"]:
        modelo._tr.getparent().remove(modelo._tr)
        _valor_celula(total, fmt_money(info["total_devolvido"]), False)


def _paragrafos_entre(cell: _Cell, ini: Paragraph, fim_padrao: str) -> List[Paragraph]:
    out, dentro = [], False
    rx = re.compile(fim_padrao)
    for q in cell.paragraphs:
        if q._p is ini._p:
            dentro = True
            continue
        if dentro:
            if rx.search(norm(q.text)):
                break
            out.append(q)
    return out


def _irregularidades_formais(info: Dict, regime: Optional[str]) -> List[Parte]:
    linhas = []
    for o in info["ocorrencias"]:
        if o["categoria"] != "a":
            continue
        if o["titulo"] == TITULO_ATRASO:
            linhas.append((ATRASO_PC.get(regime) or " OU ".join(ATRASO_PC.values()), regime not in ATRASO_PC))
        else:
            linhas.append((f"{o['titulo']}: {_sem_pendente(o['descricao'])}", True))
    return linhas


def _inconsistencias_financeiras(info: Dict) -> List[Parte]:
    ocorr = [o for o in info["ocorrencias"] if o["categoria"] in ("c", "e", "f")]
    return [(t, True) for t, _ in _linhas_ocorrencias(ocorr, compacto=True)]


def _substituir_lista(cell: _Cell, cab_padrao: str, fim_padrao: str, linhas: List[Parte], vazio: str):
    cab = next((q for q in cell.paragraphs if re.search(cab_padrao, norm(q.text))), None)
    if cab is None:
        return
    exemplos = [q for q in _paragrafos_entre(cell, cab, fim_padrao) if q.text.strip()]
    modelo = exemplos[0] if exemplos else None
    if linhas:
        _lista_apos(cab, linhas, modelo)
    else:
        _set_text(cab, vazio, True)
    for q in exemplos:
        _remover(q)


def _preencher_conclusao(doc, info: Dict, modelo: str):
    cell, _ = _achar(doc, r"^APOS ANALISE DA PRESTACAO DE CONTAS")
    if cell is None:
        return
    regime = modelo if modelo in ATRASO_PC else info["regime_codigo"]
    df = info["df"]
    final = (df or {}).get("final")
    e = (info["parecer"] or {}).get("extracao") or {}
    glosas = _valor(e, "glosas_ou_irregularidades") or _valor(e, "valor_glosado")
    num_parecer = _valor(e, "numero")
    fl_parecer, _ = fls((info["parecer"] or {}).get("localizacao"))

    if modelo == "INEXECUCAO":
        _preencher_conclusao_inexecucao(cell, info, regime, fl_parecer)
        return

    cabecalhos = [q for q in cell.paragraphs if norm(q.text).startswith("DECRETO 4")]
    if len(cabecalhos) == 3:
        if final is not None and final > 0.005:
            escolhido = 0 if glosas else 1
        else:
            escolhido = 2
        paras = cell.paragraphs
        idx = [next(i for i, q in enumerate(paras) if q._p is c._p) for c in cabecalhos] + [len(paras)]
        for b in range(3):
            if b == escolhido:
                _remover(paras[idx[b]])
                continue
            for q in paras[idx[b]:idx[b + 1]]:
                _remover(q)

    obra = [q for q in cell.paragraphs if norm(q.text).startswith("A OBRA ATENDE")]
    if len(obra) == 2:
        manter, tirar = (obra[1], obra[0]) if glosas else (obra[0], obra[1])
        ou = _depois(cell, obra[0], r"^OU$")
        _remover(tirar)
        if ou is not None:
            _remover(ou)
        obra = [manter]
    for q in obra:
        _sub_texto(q, r"Parecer Técnico nº\s*\S+\s*de fls\.\s*\([^)]*\)",
                   f"Parecer Técnico nº {num_parecer or 'XX/XXXX'} de fls. ({fl_parecer or 'XXX-XXX'})", True)

    _substituir_lista(cell, r"^FORAM IDENTIFICADAS AS SEGUINTES IRREGULARIDADES FORMAIS",
                      r"^FORAM IDENTIFICADAS INCONSISTENCIAS FINANCEIRAS", _irregularidades_formais(info, regime),
                      "Não foram identificadas irregularidades formais nos documentos analisados.")
    _substituir_lista(cell, r"^FORAM IDENTIFICADAS INCONSISTENCIAS FINANCEIRAS",
                      r"^AS INCONSISTENCIAS|^CONSIDERANDO QUE O CONVENENTE|^CONCLUIDA A ANALISE",
                      _inconsistencias_financeiras(info),
                      "Não foram identificadas inconsistências financeiras nos documentos analisados.")

    if final is not None:
        valor = fmt_money(abs(final))
        extenso = valor_por_extenso(final)
        aviso = " [VALOR PRELIMINAR – depende da conciliação financeira]" if info["df_incerto"] else ""
        for q in cell.paragraphs:
            t = norm(q.text)
            if t.startswith("AS INCONSISTENCIAS") or t.startswith("CONSIDERANDO QUE O CONVENENTE") or t.startswith("(X"):
                novo = re.sub(r"R\$\s?X{3,}", valor + (aviso if "(X" not in q.text.upper() else ""), q.text)
                novo = re.sub(r"\(X[X\s]{3,}\)", f"({extenso}){'' if t.startswith('(X') else aviso}", novo)
                if novo != q.text:
                    _set_text(q, novo, True)
    for q in cell.paragraphs:
        if norm(q.text).startswith("ESTE VALOR, DEVIDAMENTE ATUALIZADO"):
            _sub_texto(q, r"R\$\s?[\dX][\d.,X]*\s*\([^)]*\)", "R$ [VALOR ATUALIZADO] ([valor por extenso])", True)
            _sub_texto(q, r"fls\.\s*\(\d+\)", "fls. (XXXX)", True)
    _preencher_interveniente(cell, info)


def _preencher_interveniente(cell: _Cell, info: Dict):
    interv = _valor(info["dados"], "interveniente")
    if not interv:
        return
    for q in cell.paragraphs:
        _sub_texto(q, r"interveniente\s+x{4,}", f"interveniente {interv}", True)


def _preencher_conclusao_inexecucao(cell: _Cell, info: Dict, regime: Optional[str], fl_parecer: str):
    base = BASE_INEXECUCAO.get(regime)
    if base:
        for q in cell.paragraphs:
            if _sub_texto(q, r"art\. 12, XIII, alínea a, do Decreto Estadual nº 43\.635, de 2003 OU .+?Resolução Conjunta 04/2015",
                          base, None):
                _sub_texto(q, r"\bo o\b", "o", None)
    for q in cell.paragraphs:
        if norm(q.text).startswith("A OBRA NAO FOI EXECUTADA"):
            _sub_texto(q, r"fls\. \([^)]*\)", f"fls. ({fl_parecer or 'XXX'})", True)
        if norm(q.text).startswith("ENCAMINHAMENTO DA PRESTACAO"):
            if info["atrasado"]:
                _set_text(q, ATRASO_PC.get(regime) or q.text, regime not in ATRASO_PC)
            else:
                cab = next((x for x in cell.paragraphs if norm(x.text).startswith("FORAM IDENTIFICADAS AS SEGUINTES")), None)
                if cab is not None:
                    _set_text(cab, "Não foram identificadas irregularidades formais nos documentos analisados.", True)
                _remover(q)

    devolveu = next((q for q in cell.paragraphs if norm(q.text).startswith("O CONVENENTE, APOS NOTIFICACAO")), None)
    nao_devolveu = next((q for q in cell.paragraphs if norm(q.text).startswith("O CONVENENTE NAO PROCEDEU")), None)
    if devolveu is None or nao_devolveu is None:
        return
    paras = cell.paragraphs
    i_dev = next(i for i, q in enumerate(paras) if q._p is devolveu._p)
    i_nao = next(i for i, q in enumerate(paras) if q._p is nao_devolveu._p)
    ou = next((q for q in paras[i_dev:i_nao] if norm(q.text) == "OU"), None)
    i_ou = next(i for i, q in enumerate(paras) if q._p is ou._p) if ou is not None else i_nao
    if info["devolucoes"]:
        fl_dev, _ = fls(info["devolucoes"][0].get("localizacao"))
        _sub_texto(devolveu, r"fls\. \([^)]*\)", f"fls. ({fl_dev or 'XXX'})", True)
        for q in paras[i_ou:]:
            _remover(q)
    else:
        for q in paras[i_dev:i_ou + 1]:
            _remover(q)
        _set_text(nao_devolveu, nao_devolveu.text, True)
    _preencher_interveniente(cell, info)


# ─────────────────────────────────────────────────────────────────────────────
# Formato do analista (a partir de resultado["parecer"])
# ─────────────────────────────────────────────────────────────────────────────

def _partes(lista: List[List[Any]]) -> List[Tuple]:
    return [tuple(x) for x in lista]


def _dados_parecer(doc, parecer: Dict[str, Any], info: Dict):
    d = parecer["dados"]

    def campo(padrao: str, rotulo: str, valor, fmt=None):
        _, p = _achar(doc, padrao)
        if p is None:
            return
        if valor in (None, ""):
            _set_runs(p, [(rotulo, False, True), (" NÃO IDENTIFICADO – conferir", True, False)])
        else:
            _set_runs(p, [(rotulo, False, True), (f" {fmt(valor) if fmt else valor}", False, False)])

    campo(r"^NUMERO DO CONVENIO", "Número do Convênio:", d.get("numero"))
    campo(r"^MUNICIPIO", "Município:", d.get("municipio"))
    campo(r"^SECRETARIA INTERVENIENTE", "Secretaria Interveniente:", d.get("interveniente"))
    campo(r"^OBJETO:", "Objeto:", d.get("objeto"))
    campo(r"^VIGENCIA:", "Vigência:", d.get("vigencia_fim"))
    campo(r"^VALOR CONCEDENTE", "Valor Concedente:", d.get("valor_concedente"), fmt_money)
    campo(r"^VALOR CONTRAPARTIDA", "Valor Contrapartida:", d.get("valor_contrapartida"), fmt_money)
    campo(r"^VALOR TOTAL PACTUADO", "Valor Total Pactuado:", d.get("valor_total"), fmt_money)

    pc = parecer["pc"]
    cell, p = _achar(doc, r"^PRESTACAO DE CONTAS FINAL")
    if p is None:
        return
    _set_runs(p, [("Prestação de Contas Final: ", False, True),
                  (f"Ofício nº {pc.get('oficio') or '---'}, datado de {pc.get('data') or '---/---/----'}.",
                   not (pc.get("oficio") and pc.get("data")), False)])
    caixa = _depois(cell, p, caixa=True)
    if caixa is not None:
        _marcar(caixa, CAIXA.get(pc.get("resultado")))
    cab = next((q for q in cell.paragraphs if norm(q.text).startswith("RESSALVAS")), None)
    variantes = [q for q in cell.paragraphs if norm(q.text).startswith("ENCAMINHAMENTO DA PRESTACAO") or norm(q.text) == "OU"]
    if cab is not None:
        if pc.get("ressalvas"):
            _set_runs(cab, [("Ressalvas:", False, True)])
            _lista_apos(cab, [(r, False) for r in pc["ressalvas"]], modelo=variantes[0] if variantes else None)
        elif pc.get("resultado") == "ATENDE":
            _remover(cab)
        else:
            item = info["itens"].get("Prazo da prestação de contas") or {}
            _set_runs(cab, [("Ressalvas: ", False, True),
                            (f"prazo pendente de validação – {_sem_pendente(item.get('observacao') or 'dados insuficientes')}", True, False)])
    for q in variantes:
        _remover(q)


def _bloco_parecer(doc, padrao: str, bloco: Dict[str, Any], info: Dict, itens: Tuple[str, ...]):
    cell, p = _achar(doc, padrao)
    if p is None:
        return
    _set_runs(p, _partes(bloco["cabecalho"]))
    resultado = bloco.get("resultado")
    caixa = _depois(cell, p, caixa=True)
    if caixa is not None:
        _marcar(caixa, CAIXA.get(resultado))
    cab = _depois(cell, p, r"^IRREGULARIDADES")
    lista = _depois(cell, cab, r"^LISTAR AS IRREGULARIDADES") if cab is not None else None
    linhas: List[Parte] = [tuple(x) for x in bloco.get("linhas") or []]
    if resultado not in CAIXA:
        achado = next((info["itens"][i] for i in itens if i in info["itens"]), None)
        obs = _sem_pendente(achado["observacao"]) if achado else "item não avaliado automaticamente"
        linhas = [(f"Pendente de validação do analista: {obs}", True)] + linhas
    if resultado == "ATENDE" or not linhas:
        for q in (lista, cab):
            if q is not None:
                _remover(q)
        return
    if cab is not None:
        _set_runs(cab, [(CABECALHO_IRREGULARIDADES.get(resultado, "Irregularidades e/ou invalidades:"), False, True)])
    if lista is not None:
        _set_runs(lista, [(linhas[0][0], linhas[0][1], False)])
        _lista_apos(lista, linhas[1:])
    elif caixa is not None:
        _lista_apos(caixa, linhas)


def _documentacao_parecer(doc, parecer: Dict[str, Any], info: Dict):
    b = parecer["blocos"]
    _bloco_parecer(doc, r"^PROCESSO LICITATORIO N", b["licitacao"], info, BLOCO_LICITACAO)
    _bloco_parecer(doc, r"^ORDEM DE SERVICO \(FLS", b["ordem_servico"], info, BLOCO_OS)
    _bloco_parecer(doc, r"^COMPROVANTES DE DESPESAS \(FLS", b["despesas"], info, BLOCO_DESPESAS)
    _bloco_parecer(doc, r"^MOVIMENTACAO BANCARIA", b["movimentacao"], info, BLOCO_BANCO)
    for padrao, itens in ((r"^PROCESSO LICITATORIO \(", BLOCO_LICITACAO), (r"^ORDEM DE SERVICO \( OBJETO", BLOCO_OS),
                          (r"^COMPROVANTES DE DESPESAS \( OBJETO", BLOCO_DESPESAS)):
        _preencher_bloco(doc, padrao, None, info, itens)


def _consideracoes_parecer(doc, info: Dict, diligencia: bool, numero: Optional[str], fl_oficio: Optional[str]):
    """Em diligência a AD só registra os blocos; observações gerais ficam no Relatório de Validação."""
    _, instr = _achar(doc, r"^QUANDO O ANALISTA CONSIDERAR")
    if instr is not None:
        _remover(instr)
    _, base = _achar(doc, r"^COM BASE N[OA]")
    if base is not None and ("XXX/XXXX" in base.text.upper() or "FLS. (XXX)" in base.text.upper()):
        texto = re.sub(r"Fls\. \(xxx\)", f"Fls. ({fl_oficio or 'xxx'})", base.text, flags=re.IGNORECASE)
        texto = re.sub(r"Convênio nº xxx/xxxx", f"Convênio nº {numero or 'xxx/xxxx'}", texto, flags=re.IGNORECASE)
        _set_text(base, texto)
    _, cab = _achar(doc, r"^CONSIDERACOES GERAIS")
    linhas = [] if diligencia else _consideracoes(info)
    if not linhas:
        if cab is not None:
            _remover(cab)
        return
    if cab is not None:
        _set_runs(cab, [("CONSIDERAÇÕES GERAIS:", False, True)])
    if base is not None:
        _lista_apos(base, linhas)


def _nota_minuta(doc, nome_dossie: str, modelo: str, sugestao: Dict[str, Any], diligencia: bool = False):
    titulo = next((p for p in doc.paragraphs if norm(p.text).startswith("ANALISE DOCUMENTAL")), None)
    if titulo is None:
        return
    texto = (f"Minuta gerada automaticamente em {datetime.now().strftime('%d/%m/%Y %H:%M')} a partir do dossiê "
             f"“{nome_dossie}” (modelo: {MODELOS_AD[modelo][1]}). Trechos destacados em amarelo dependem de conferência "
             "ou complementação pelo analista; as caixas marcadas refletem a sugestão do sistema.")
    if diligencia:
        texto += (" Há documentos a solicitar ao convenente (ver Anexo – Solicitação de Documentos): as seções III a VI "
                  "foram mantidas em branco até a resposta da diligência.")
    if modelo != sugestao["modelo_sugerido"] and modelo != "INEXECUCAO":
        texto += f" ATENÇÃO: o modelo escolhido difere do sugerido ({sugestao['motivo']})"
    elif sugestao["modelo_sugerido"] == modelo and "confirme" in sugestao["motivo"]:
        texto += f" ATENÇÃO: {sugestao['motivo']}"
    novo = deepcopy(titulo._p)
    titulo._p.addnext(novo)
    par = Paragraph(novo, titulo._parent)
    _set_text(par, texto, True)
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in par.runs:
        r.bold = False
        r.italic = True
        r.font.size = Pt(8)


def gerar_minuta_ad(resultado: Dict[str, Any], nome_dossie: str, modelo: Optional[str] = None) -> Tuple[bytes, str]:
    sugestao = sugerir_modelo(resultado)
    modelo = (modelo or sugestao["modelo_sugerido"]).upper()
    if modelo not in MODELOS_AD:
        raise ValueError(f"Modelo inválido: {modelo}. Use um de: {', '.join(MODELOS_AD)}")
    info = coletar_dados(resultado)
    doc = Document(str(MODELOS_DIR / MODELOS_AD[modelo][0]))
    numero = _valor(info["dados"], "numero_convenio")
    parecer = resultado.get("parecer") if modelo != "INEXECUCAO" else None

    if parecer:
        numero = parecer["dados"].get("numero") or numero
        diligencia = bool(parecer.get("diligencia"))
        _dados_parecer(doc, parecer, info)
        _consideracoes_parecer(doc, info, diligencia, numero, parecer["pc"].get("fls"))
        _documentacao_parecer(doc, parecer, info)
    else:
        diligencia = False
        _preencher_dados(doc, info, modelo)
        _preencher_consideracoes(doc, info, numero)
        _preencher_documentacao(doc, info)
    if not diligencia:
        _preencher_parecer(doc, info)
        _preencher_gecon(doc, info)
        _preencher_df(doc, info)
        _preencher_devolucoes(doc, info)
        _preencher_conclusao(doc, info, modelo)
    _nota_minuta(doc, nome_dossie, modelo, sugestao, diligencia)

    buf = io.BytesIO()
    doc.save(buf)
    rotulo = re.sub(r"[^\w.-]+", "_", f"AD_CV_{numero or 'sem_numero'}_{MODELOS_AD[modelo][1]}").strip("_")
    return buf.getvalue(), f"{rotulo}.docx"
